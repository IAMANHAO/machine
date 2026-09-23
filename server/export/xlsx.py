"""Excel 导出。

与 PDF 同源（都从 report.Report 投影），所以两份文件里的每一格必然一致。
Excel 面向的是"拿去接着算"的场景，所以数值列尽量写成真数字而不是文本。
"""

from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .report import Report, Table

INK = "1F2933"
MUTED = "6B7A88"
HEAD_BG = "EEF3F7"
OK = "0F7A4F"
WARN = "A15C00"
ERR = "B4232A"

_THIN = Side(style="thin", color="D4DCE3")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
WRAP = Alignment(vertical="top", wrap_text=True)

_NUMERIC = re.compile(r"^-?\d+(?:\.\d+)?$")


def _maybe_number(text: str):
    """纯数字写成数字，带单位的保持文本——别把 '76.2 mm' 硬塞成数字。"""
    t = str(text).strip()
    return float(t) if _NUMERIC.match(t) else text


def render(rep: Report) -> bytes:
    wb = Workbook()
    _cover(wb.active, rep)

    for tbl in rep.tables:
        ws = wb.create_sheet(_sheet_name(tbl.title))
        _sheet(ws, tbl)

    if rep.procure_keyword:
        _procure(wb.create_sheet("采购信息"), rep)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet_name(title: str) -> str:
    # Excel 工作表名不能超 31 字符，也不能含 []:*?/\
    name = re.sub(r"[\[\]:*?/\\]", "", title.replace(" · ", " "))
    return name[:31]


def _cover(ws, rep: Report) -> None:
    ws.title = "概要"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 96

    ws["A1"] = rep.title
    ws["A1"].font = Font(size=15, bold=True, color=INK)
    ws.merge_cells("A1:B1")

    tone = {"ok": OK, "check_failed": ERR, "data_missing": ERR,
            "no_solution": WARN, "needs_choice": WARN}.get(rep.status, MUTED)

    rows = [
        ("依据标准", rep.standard),
        ("物料标识", rep.material),
        ("生成时间", rep.generated_at),
        ("结论", rep.status_label),
        ("数据核验状态", rep.confidence_label),
    ]
    r = 3
    for label, value in rows:
        ws.cell(r, 1, label).font = Font(bold=True, color=MUTED)
        c = ws.cell(r, 2, value)
        c.alignment = WRAP
        if label == "结论":
            c.font = Font(bold=True, color=tone)
        elif label == "数据核验状态" and rep.confidence != "verified":
            c.font = Font(color=WARN)
        r += 1

    r += 1
    if rep.blocker:
        ws.cell(r, 1, "流程未走完").font = Font(bold=True, color=ERR)
        ws.cell(r, 2, rep.blocker).alignment = WRAP
        r += 1
    for w in rep.warnings:
        ws.cell(r, 1, "风险提示").font = Font(bold=True, color=WARN)
        ws.cell(r, 2, w).alignment = WRAP
        ws.row_dimensions[r].height = 30
        r += 1

    r += 1
    ws.cell(r, 1, "说明").font = Font(bold=True, color=MUTED)
    c = ws.cell(r, 2, rep.disclaimer)
    c.alignment = WRAP
    ws.row_dimensions[r].height = 30

    ws.sheet_view.showGridLines = False


def _sheet(ws, tbl: Table) -> None:
    ws.cell(1, 1, tbl.title).font = Font(size=12, bold=True, color=INK)

    head_row = 3
    for j, col in enumerate(tbl.columns, start=1):
        c = ws.cell(head_row, j, col)
        c.font = Font(bold=True, color=INK)
        c.fill = PatternFill("solid", fgColor=HEAD_BG)
        c.border = BORDER
        c.alignment = Alignment(vertical="center", wrap_text=True)

    concl = tbl.columns.index("结论") + 1 if "结论" in tbl.columns else None

    for i, row in enumerate(tbl.rows, start=head_row + 1):
        for j, val in enumerate(row, start=1):
            c = ws.cell(i, j, _maybe_number(val))
            c.border = BORDER
            c.alignment = WRAP
            if concl and j == concl and str(val) != "通过":
                c.font = Font(bold=True, color=ERR)

    if tbl.note:
        nr = head_row + len(tbl.rows) + 2
        c = ws.cell(nr, 1, tbl.note)
        c.font = Font(size=9, color=MUTED)
        c.alignment = WRAP
        ws.merge_cells(start_row=nr, start_column=1,
                       end_row=nr, end_column=max(len(tbl.columns), 2))

    _autosize(ws, tbl)
    ws.freeze_panes = ws.cell(head_row + 1, 1)
    ws.sheet_view.showGridLines = False


def _autosize(ws, tbl: Table) -> None:
    """按内容估宽。中文按两个字符宽算，否则列会挤在一起。"""
    for j, col in enumerate(tbl.columns, start=1):
        longest = _width(col)
        for row in tbl.rows:
            if j - 1 < len(row):
                for line in str(row[j - 1]).split("\n"):
                    longest = max(longest, _width(line))
        ws.column_dimensions[get_column_letter(j)].width = min(max(longest + 3, 9), 48)


def _width(text: str) -> int:
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in str(text))


def _procure(ws, rep: Report) -> None:
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 96

    ws.cell(1, 1, "采购信息").font = Font(size=12, bold=True, color=INK)
    ws.cell(3, 1, "采购关键词").font = Font(bold=True, color=MUTED)
    ws.cell(3, 2, rep.procure_keyword).font = Font(bold=True)

    r = 5
    ws.cell(r, 1, "渠道").font = Font(bold=True, color=INK)
    ws.cell(r, 2, "搜索链接").font = Font(bold=True, color=INK)
    for c in (ws.cell(r, 1), ws.cell(r, 2)):
        c.fill = PatternFill("solid", fgColor=HEAD_BG)
        c.border = BORDER
    r += 1

    for name, url in rep.procure_links:
        ws.cell(r, 1, name).border = BORDER
        c = ws.cell(r, 2, url)
        c.hyperlink = url
        c.font = Font(color="1C5FA8", underline="single")
        c.border = BORDER
        r += 1

    if rep.procure_note:
        r += 1
        c = ws.cell(r, 1, rep.procure_note)
        c.font = Font(size=9, color=MUTED)
        c.alignment = WRAP
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.row_dimensions[r].height = 42

    ws.sheet_view.showGridLines = False
