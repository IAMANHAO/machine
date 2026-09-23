"""导出测试（M5）。

导出件是要拿去评审、归档、当采购依据的，所以它比界面更不能有第二套说法。
重点守三件事：

1. **信源清单必须在**，且如实反映核验状态——这是本产品的核心承诺
2. **代入式必须在**，"每一次选型都能被逐行验证"在离线纸面上同样成立
3. **流程没走完时不能假装走完了**
"""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

GOLDEN = {
    "P": 5.5, "n1": 1450, "i": 2, "a0": 400, "belt_type": "H",
    "prime_mover": "ac_motor_normal", "work_machine": "medium_uniform",
    "hours_per_day": "h_le_10",
}


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("MDS_DATA_DIR", str(tmp_path / "data"))
    from server.main import app
    with TestClient(app) as c:
        yield c


def _export(client, fmt: str, **overrides):
    values = {**GOLDEN, **overrides}
    for k in [k for k, v in values.items() if v is None]:
        values.pop(k)
    return client.post(f"/api/export/{fmt}",
                       json={"material": "synchronous_belt", "values": values})


def _pdf_text(content: bytes) -> str:
    return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(content)).pages)


# --- 通用 ------------------------------------------------------------------

@pytest.mark.parametrize("fmt,magic", [("pdf", b"%PDF"), ("xlsx", b"PK")])
def test_export_returns_a_real_file(client, fmt, magic):
    r = _export(client, fmt)
    assert r.status_code == 200
    assert r.content.startswith(magic)
    assert len(r.content) > 5000
    assert r.headers["x-report-status"] == "ok"
    assert r.headers["x-report-confidence"] == "single_source"


@pytest.mark.parametrize("fmt", ["pdf", "xlsx"])
def test_chinese_filename_uses_rfc5987(client, fmt):
    """不走 RFC 5987 的话中文文件名在浏览器里会变成乱码。"""
    cd = _export(client, fmt).headers["content-disposition"]
    assert "filename*=UTF-8''" in cd
    assert "%E9%80%89%E5%9E%8B%E6%8A%A5%E5%91%8A" in cd  # 选型报告


def test_unknown_format_is_404(client):
    assert client.post("/api/export/docx", json={
        "material": "synchronous_belt", "values": GOLDEN}).status_code == 404


def test_invalid_input_is_422_not_a_broken_file(client):
    r = _export(client, "pdf", P=99999)
    assert r.status_code == 422
    assert r.json()["detail"]["param"] == "P"


# --- PDF -------------------------------------------------------------------

def test_pdf_carries_the_full_derivation(client):
    txt = _pdf_text(_export(client, "pdf").content)
    assert "表 1 · 计算过程汇总" in txt
    assert "1.1 × 5.5" in txt, "代入式必须在纸面上，否则无从逐行复核"
    assert "6.05" in txt
    assert "GB/T 11362-2021" in txt


def test_pdf_carries_the_source_list_with_real_status(client):
    txt = _pdf_text(_export(client, "pdf").content)
    assert "表 4 · 信源清单" in txt
    assert "condition_factors.yaml" in txt
    assert "rated_power.yaml" in txt
    assert "单一信源" in txt, "未核验状态必须如实写在报告里"
    assert "正式投产前请对照标准原件复核" in txt


def test_pdf_states_ai_did_not_compute(client):
    txt = _pdf_text(_export(client, "pdf").content)
    assert "AI 不参与任何数值计算" in txt


def test_pdf_translates_enum_values(client):
    """报告里留一串 ac_motor_normal，对着图纸评审的人是没法读的。"""
    txt = _pdf_text(_export(client, "pdf").content)
    assert "交流电动机（启动转矩正常）" in txt
    assert "中等载荷（一般机床、搅拌器）" in txt


def test_pdf_has_no_emoji(client):
    """CID 字体没有 emoji 字形，留着只会变成空白方框。"""
    txt = _pdf_text(_export(client, "pdf").content)
    assert not [ch for ch in txt if ord(ch) > 0x1F000]


def test_pdf_includes_procurement(client):
    txt = _pdf_text(_export(client, "pdf").content)
    assert "采购关键词" in txt
    assert "s.taobao.com" in txt


# --- 流程未走完 ------------------------------------------------------------

def test_export_of_a_blocked_run_says_so(client):
    """n1=4000 数据缺失。导出件不能假装选型成功了。"""
    r = _export(client, "pdf", n1=4000)
    assert r.status_code == 200
    assert r.headers["x-report-status"] == "data_missing"
    txt = _pdf_text(r.content)
    assert "数据缺失" in txt
    assert "流程未走完" in txt
    assert "表 3 · 最终选型结果" not in txt, "没有结果就不该有结果表"


def test_export_marks_unexecuted_checks(client):
    txt = _pdf_text(_export(client, "pdf", n1=4000).content)
    assert "未执行" in txt
    assert "未执行不等于通过" in txt


def test_export_of_no_solution_distinguishes_from_missing_data(client):
    r = _export(client, "pdf", belt_type="L")
    assert r.headers["x-report-status"] == "no_solution"
    txt = _pdf_text(r.content)
    assert "该规格系列内无解" in txt


# --- Excel -----------------------------------------------------------------

def test_xlsx_sheets_cover_everything(client):
    wb = load_workbook(io.BytesIO(_export(client, "xlsx").content))
    names = wb.sheetnames
    assert names[0] == "概要"
    assert any("计算过程汇总" in n for n in names)
    assert any("信源清单" in n for n in names)
    assert any("最终选型结果" in n for n in names)
    assert "采购信息" in names


def test_xlsx_numbers_are_numbers(client):
    """Excel 面向"拿去接着算"，纯数值列不该是文本。"""
    wb = load_workbook(io.BytesIO(_export(client, "xlsx").content))
    ws = next(wb[n] for n in wb.sheetnames if "计算过程汇总" in n)
    seq = [ws.cell(r, 1).value for r in range(4, 4 + 5)]
    assert seq == [1.0, 2.0, 3.0, 4.0, 5.0], "序号列应当是数字"


def test_xlsx_source_sheet_has_verification_columns(client):
    wb = load_workbook(io.BytesIO(_export(client, "xlsx").content))
    ws = next(wb[n] for n in wb.sheetnames if "信源清单" in n)
    header = [ws.cell(3, c).value for c in range(1, 7)]
    assert header == ["数据表", "主信源", "第二信源", "核验状态", "核验日期", "内容指纹"]
    body = "\n".join(str(ws.cell(r, c).value) for r in range(4, 10) for c in range(1, 7))
    assert "单一信源" in body
    assert "（无）" in body, "没有第二信源要如实写"


def test_xlsx_procurement_links_are_clickable(client):
    wb = load_workbook(io.BytesIO(_export(client, "xlsx").content))
    ws = wb["采购信息"]
    assert ws.cell(6, 2).hyperlink is not None
    assert str(ws.cell(6, 2).value).startswith("https://")


# --- 与引擎一致 ------------------------------------------------------------

def test_export_matches_what_the_engine_computed(client):
    """导出件不是界面的打印稿——服务端自己重跑引擎，两边必须对得上。"""
    run = client.post("/api/selection/run", json={
        "material": "synchronous_belt", "values": GOLDEN, "save": False}).json()
    bs = run["trace"]["outputs"]["bs"]

    txt = _pdf_text(_export(client, "pdf").content)
    assert str(bs) in txt

    wb = load_workbook(io.BytesIO(_export(client, "xlsx").content))
    ws = next(wb[n] for n in wb.sheetnames if "最终选型结果" in n)
    body = "\n".join(str(ws.cell(r, 2).value) for r in range(4, 20))
    assert str(bs) in body


LUBRICANT_CHAIN = {
    "target": "chain", "T_work": 55, "z1": 19, "p_chain": 12.7,
    "n_chain": 300, "chain_method": "hand",
}


def test_branch_workflow_exports(client):
    """分支型工作流同样要能出报告。"""
    r = client.post("/api/export/pdf", json={
        "material": "lubricant", "values": LUBRICANT_CHAIN})
    assert r.status_code == 200
    txt = _pdf_text(r.content)
    assert "链条" in txt
    assert "150 / 220 / 320" in txt, "推荐区间要完整列出，不能只给一个值"
    assert "chain_viscosity.yaml" in txt


def test_export_omits_label_steps_and_other_branches(client):
    """表 1 是计算过程汇总：凑标签的 text 步骤与别的分支都不该出现在里面。"""
    txt = _pdf_text(client.post("/api/export/pdf", json={
        "material": "lubricant", "values": LUBRICANT_CHAIN}).content)
    assert "{v_chain}" not in txt, "未渲染的模板不能漏进报告"
    assert "对象标签" not in txt
    assert "节圆速" not in txt, "开式齿轮的步骤不属于链条这条流程"
